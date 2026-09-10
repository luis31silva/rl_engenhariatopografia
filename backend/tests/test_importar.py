"""Testes do importador do .xlsm (estrutura trabalho + itens)."""
from datetime import date, datetime
from decimal import Decimal

from openpyxl import Workbook

from app import importar as imp


# ---- Funções de conversão ----
def test_texto_trata_vazios():
    assert imp._texto("-") is None
    assert imp._texto("---") is None
    assert imp._texto("") is None
    assert imp._texto(None) is None
    assert imp._texto("  Almerindo  ") == "Almerindo"


def test_decimal_valido_e_invalido():
    assert imp._decimal("380") == Decimal("380")
    assert imp._decimal(150.5) == Decimal("150.5")
    assert imp._decimal("-") is None
    assert imp._decimal(None) is None
    assert imp._decimal("abc") is None


def test_data_datetime_e_serial():
    assert imp._data(datetime(2015, 1, 1)) == date(2015, 1, 1)
    assert imp._data("-") is None
    assert imp._data(42005) == date(2015, 1, 1)


# ---- Integração: workbook em memória ----
def _construir_workbook(prefixo: str = "IMP") -> Workbook:
    wb = Workbook()

    dados = wb.active
    dados.title = "Dados"
    for _ in range(20):
        dados.append([None] * 16)
    dados.cell(row=1, column=1, value="INTERMEDIÁRIOS")
    dados.cell(row=1, column=6, value="EXTERNOS")
    dados.cell(row=3, column=1, value=f"{prefixo} Arqt Gil")
    dados.cell(row=3, column=3, value=1030)
    dados.cell(row=3, column=6, value=f"{prefixo} Nicola")
    dados.cell(row=3, column=15, value=f"{prefixo} Marta")  # externo extra
    # Especialidades (col O=15, L=12 top, M=13 ok), a partir da linha 18.
    dados.cell(row=18, column=15, value=f"{prefixo}-LEVTOP")
    dados.cell(row=18, column=12, value=5)
    dados.cell(row=18, column=13, value=7)
    dados.cell(row=19, column=15, value=f"{prefixo}-ENG")
    dados.cell(row=19, column=12, value=30)
    dados.cell(row=19, column=13, value=40)

    trab = wb.create_sheet("Trabalhos")
    trab.append([
        "PONT.", "REF.", "CLIENTE", "CONTATO", "LOCALIDADE", "ENDEREÇO",
        "INTERMEDIÁRIO", "ESPEC.", "VALOR", "ADJUDC.", "ENTREGA", "EXT.",
        "VALOR", "AVAL.", "€NG", "DES./OBS.",
    ])
    # Trabalho A: linha inicial (LEVTOP, situação "Pago") + continuação (ENG, sem situação -> herda).
    trab.append([
        "Pago", f"{prefixo}-001.15", f"{prefixo} Venceslau", "-", "Agrela", "Rua X",
        f"{prefixo} Arqt Gil", f"{prefixo}-LEVTOP", 200, datetime(2023, 10, 18), datetime(2023, 10, 20),
        "-", 0, "TOP", None, "obs",
    ])
    trab.append([
        None, None, None, None, None, None,
        None, f"{prefixo}-ENG", 2275.5, datetime(2025, 2, 17), datetime(2025, 5, 21),
        f"{prefixo} Nicola", 300, "MAU", None, None,
    ])
    # Trabalho B: uma só especialidade, situação "Falta pagamento".
    trab.append([
        "Falta pagamento", f"{prefixo}-002.15", f"{prefixo} Elisabete", None, None, None,
        "-", f"{prefixo}-LEVTOP", 150, datetime(2015, 1, 2), None,
        "-", 0, None, None, None,
    ])
    # Linha totalmente vazia.
    trab.append([None] * 16)

    # Aba "Relatório Contas" com 2 blocos de equipamento (o 2º herda a %).
    rc = wb.create_sheet("Relatório Contas")
    rc.append(["DEPRECIAÇÃO"])  # linha 1
    # linha 2: nomes nos cols 1 e 4.
    l2 = [None] * 8
    l2[1] = f"{prefixo} Estação"
    l2[4] = f"{prefixo} GPS"
    rc.append(l2)
    # linha 3: cabeçalhos (irrelevante para o parsing, mas mantém a estrutura).
    rc.append([None] * 8)
    # linha 4: dados. col0=% (0.1666), col1=ano, col2=valor | col4=ano, col5=valor (herda %).
    l4 = [None] * 8
    l4[0] = 0.1666
    l4[1] = 2021
    l4[2] = 23000
    l4[4] = 2020
    l4[5] = 5000
    rc.append(l4)

    return wb


def test_importacao_agrupa_itens_por_trabalho(client, monkeypatch, tmp_path):
    from app.database import SessionLocal

    monkeypatch.setattr(imp, "SessionLocal", SessionLocal)
    wb = _construir_workbook(prefixo="INTG")
    ficheiro = tmp_path / "teste.xlsx"
    wb.save(ficheiro)

    rel = imp.importar(str(ficheiro))
    assert rel.trabalhos_criados == 2
    assert rel.itens_criados == 3  # A: 2 itens, B: 1 item

    # Trabalho A deve ter 2 itens (LEVTOP + ENG).
    resp = client.get("/trabalhos", params={"q": "INTG-001.15"})
    items = resp.json()["items"]
    a = next(t for t in items if t["referencia"] == "INTG-001.15")
    assert a["cliente"] == "INTG Venceslau"
    assert len(a["itens"]) == 2
    codigos = {i["especialidade_codigo"] for i in a["itens"]}
    assert codigos == {"INTG-LEVTOP", "INTG-ENG"}
    # valor_total = 200 + 2275.5
    assert float(a["valor_total"]) == 2475.5
    # O item ENG tem externo associado.
    eng = next(i for i in a["itens"] if i["especialidade_codigo"] == "INTG-ENG")
    assert eng["externo_nome"] == "INTG Nicola"
    # Situação: a linha inicial era "Pago"; o ENG (vazio) herda "Pago".
    assert all(i["situacao"] == "Pago" for i in a["itens"])
    assert a["situacao_trabalho"] == "Pago"


def test_importacao_idempotente(client, monkeypatch, tmp_path):
    from app.database import SessionLocal

    monkeypatch.setattr(imp, "SessionLocal", SessionLocal)
    wb = _construir_workbook(prefixo="IDEM")
    ficheiro = tmp_path / "teste2.xlsx"
    wb.save(ficheiro)

    imp.importar(str(ficheiro))
    rel2 = imp.importar(str(ficheiro))
    assert rel2.trabalhos_criados == 0
    assert rel2.trabalhos_ignorados >= 2


def test_situacao_pela_cor_de_fundo(client, monkeypatch, tmp_path):
    """A situação vem da COR da coluna A (verde=Pago), mesmo sem texto."""
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    from app.database import SessionLocal

    monkeypatch.setattr(imp, "SessionLocal", SessionLocal)

    wb = Workbook()
    dados = wb.active
    dados.title = "Dados"
    dados.append(["INTERMEDIÁRIOS"])

    trab = wb.create_sheet("Trabalhos")
    trab.append([
        "PONT.", "REF.", "CLIENTE", "CONTATO", "LOCALIDADE", "ENDEREÇO",
        "INTERMEDIÁRIO", "ESPEC.", "VALOR", "ADJUDC.", "ENTREGA", "EXT.",
        "VALOR", "AVAL.", "€NG", "DES./OBS.",
    ])
    # Célula A sem texto, mas pintada de verde (92D050) -> deve ficar "Pago".
    trab.append([
        None, "COR-001.15", "Cliente Cor", None, None, None,
        None, None, 100, None, None, None, None, None, None, None,
    ])
    verde = PatternFill(start_color="FF92D050", end_color="FF92D050", fill_type="solid")
    trab.cell(row=2, column=1).fill = verde

    ficheiro = tmp_path / "cor.xlsx"
    wb.save(ficheiro)

    imp.importar(str(ficheiro))
    resp = client.get("/trabalhos", params={"q": "COR-001.15"})
    t = next(x for x in resp.json()["items"] if x["referencia"] == "COR-001.15")
    assert t["situacao_trabalho"] == "Pago"
    assert t["itens"][0]["situacao"] == "Pago"


def test_importa_equipamentos_com_percentagem_herdada(client, monkeypatch, tmp_path):
    from app.database import SessionLocal

    monkeypatch.setattr(imp, "SessionLocal", SessionLocal)
    wb = _construir_workbook(prefixo="EQP")
    ficheiro = tmp_path / "eq.xlsx"
    wb.save(ficheiro)

    rel = imp.importar(str(ficheiro))
    assert rel.equipamentos == 2

    # Verifica via API (endpoint de equipamentos do dashboard).
    resp = client.get("/dashboard/equipamentos")
    nomes = {e["nome"]: e for e in resp.json()}
    est = nomes.get("EQP Estação")
    gps = nomes.get("EQP GPS")
    assert est is not None and gps is not None
    assert est["ano_aquisicao"] == 2021 and float(est["valor"]) == 23000
    # O GPS herda a percentagem 0.1666 do bloco anterior.
    assert abs(float(gps["percentagem"]) - 0.1666) < 1e-6
    assert gps["ano_aquisicao"] == 2020 and float(gps["valor"]) == 5000


def test_referencia_duplicada_no_ficheiro_e_renumerada(client, monkeypatch, tmp_path):
    """Duas linhas com a mesma referência no ficheiro: a 2ª recebe novo número."""
    from openpyxl import Workbook
    from app.database import SessionLocal

    monkeypatch.setattr(imp, "SessionLocal", SessionLocal)

    wb = Workbook()
    dados = wb.active
    dados.title = "Dados"
    dados.append(["INTERMEDIÁRIOS"])

    trab = wb.create_sheet("Trabalhos")
    trab.append([
        "PONT.", "REF.", "CLIENTE", "CONTATO", "LOCALIDADE", "ENDEREÇO",
        "INTERMEDIÁRIO", "ESPEC.", "VALOR", "ADJUDC.", "ENTREGA", "EXT.",
        "VALOR", "AVAL.", "€NG", "DES./OBS.",
    ])
    # Dois trabalhos com a MESMA referência DUP-092.20 e um seguinte.
    trab.append([None, "DUP-092.20", "Cliente A", None, None, None, None, None, 100, None, None, None, None, None, None, None])
    trab.append([None, "DUP-092.20", "Cliente B", None, None, None, None, None, 200, None, None, None, None, None, None, None])
    ficheiro = tmp_path / "dup.xlsx"
    wb.save(ficheiro)

    rel = imp.importar(str(ficheiro))
    # Ambos os trabalhos são criados (nenhum perdido).
    assert rel.trabalhos_criados == 2
    assert len(rel.referencias_renumeradas) == 1
    antiga, nova = rel.referencias_renumeradas[0]
    assert antiga == "DUP-092.20"
    assert nova != "DUP-092.20"

    # Confirma via API que existem dois trabalhos distintos, um com a nova referência.
    resp = client.get("/trabalhos", params={"q": "Cliente B"})
    b = next(t for t in resp.json()["items"] if t["cliente"] == "Cliente B")
    assert b["referencia"] == nova
