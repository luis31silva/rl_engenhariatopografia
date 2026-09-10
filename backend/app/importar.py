"""Importa os dados históricos da folha de cálculo (.xlsm) para a base de dados.

Estrutura da aba "Trabalhos": cada trabalho ocupa uma ou mais linhas consecutivas.
- A linha que inicia um trabalho tem a REFERÊNCIA preenchida (e normalmente o cliente).
- As linhas seguintes SEM referência pertencem ao mesmo trabalho.
- Cada linha (incluindo a primeira) corresponde a uma especialidade (um item), com o
  seu valor, datas de adjudicação/entrega, externo e valor pago ao externo.

Usa os modelos SQLAlchemy, pelo que funciona em SQLite (local) ou MySQL (produção),
consoante a DATABASE_URL.

Uso:
    python -m app.importar "../CONT- RLENGTOP.xlsm"
    DATABASE_URL="mysql+pymysql://..." python -m app.importar "caminho.xlsm"

Requer que as tabelas já existam (correr `alembic upgrade head` antes).
Idempotente: um trabalho já existente (mesma referência) é ignorado.
"""
from __future__ import annotations

import sys
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.database import SessionLocal
from app.models.apoio import Especialidade, Externo, Intermediario
from app.models.financeiro import Equipamento
from app.models.trabalho import Trabalho, TrabalhoItem
from app.services.situacao import Situacao

# ---- Índices de coluna (0-based) ----
# Aba "Dados"
DADOS_INTERM_NOME = 0
DADOS_INTERM_VALOR = 2
DADOS_EXTERNO_NOME = 5
DADOS_ESPEC_PRAZO_TOP = 11  # coluna L
DADOS_ESPEC_PRAZO_OK = 12   # coluna M
DADOS_ESPEC_CODIGO = 14     # coluna O
DADOS_ESPEC_PRIMEIRA_LINHA = 18

# Aba "Trabalhos"
T_SITUACAO = 0  # coluna A "PONT. SITUAÇÃO"
T_REF = 1
T_CLIENTE = 2
T_CONTACTO = 3
T_LOCALIDADE = 4
T_ENDERECO = 5
T_INTERMEDIARIO = 6
T_ESPECIALIDADE = 7
T_VALOR = 8
T_ADJUDICACAO = 9
T_ENTREGA = 10
T_EXTERNO = 11
T_VALOR_EXTERNO = 12
T_OBSERVACOES = 15

# Aba "Relatório Contas" — blocos de equipamento (depreciação).
# Cada bloco: nome (linha 2), e na linha 4 (primeira de dados) percentagem/ano/valor.
# `col_pct=None` significa que herda a percentagem do bloco anterior (16,66%).
RC_EQUIP_LINHA_NOME = 2
RC_EQUIP_LINHA_DADOS = 4
RC_EQUIPAMENTOS = [
    {"col_nome": 1, "col_pct": 0, "col_ano": 1, "col_valor": 2},
    {"col_nome": 4, "col_pct": None, "col_ano": 4, "col_valor": 5},
    {"col_nome": 7, "col_pct": None, "col_ano": 7, "col_valor": 8},
    {"col_nome": 11, "col_pct": 10, "col_ano": 11, "col_valor": 12},
    {"col_nome": 15, "col_pct": 14, "col_ano": 15, "col_valor": 16},
]

VAZIOS = {None, "", "-", "---", "--"}


@dataclass
class Relatorio:
    especialidades: int = 0
    intermediarios: int = 0
    externos: int = 0
    equipamentos: int = 0
    trabalhos_criados: int = 0
    itens_criados: int = 0
    trabalhos_ignorados: int = 0
    linhas_vazias: int = 0
    referencias_renumeradas: list[tuple[str, str]] = field(default_factory=list)
    nomes_nao_resolvidos: set[str] = field(default_factory=set)

    def imprimir(self) -> None:
        print("\n===== Relatório de importação =====")
        print(f"  Especialidades:        {self.especialidades}")
        print(f"  Intermediários:        {self.intermediarios}")
        print(f"  Externos:              {self.externos}")
        print(f"  Equipamentos:          {self.equipamentos}")
        print(f"  Trabalhos criados:     {self.trabalhos_criados}")
        print(f"  Itens (especialidades): {self.itens_criados}")
        print(f"  Trabalhos ignorados:   {self.trabalhos_ignorados} (referência já existia na BD)")
        print(f"  Linhas em branco:      {self.linhas_vazias}")
        if self.referencias_renumeradas:
            print(f"  Referências duplicadas renumeradas: {len(self.referencias_renumeradas)}")
            for antiga, nova in self.referencias_renumeradas:
                print(f"      {antiga} -> {nova}")
        if self.nomes_nao_resolvidos:
            print(f"  Referências de apoio criadas ao voo: {len(self.nomes_nao_resolvidos)}")


# ---- Helpers ----
def _texto(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return None if s in VAZIOS else s


def _decimal(v) -> Decimal | None:
    if v in VAZIOS:
        return None
    try:
        d = Decimal(str(v))
        return d if d >= 0 else None
    except (InvalidOperation, ValueError):
        return None


def _data(v) -> date | None:
    if v in VAZIOS:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        from datetime import timedelta

        return (datetime(1899, 12, 30) + timedelta(days=int(v))).date()
    except (ValueError, TypeError):
        return None


def _inteiro(v) -> int:
    try:
        return int(v)
    except (ValueError, TypeError):
        return 0


def _linha_totalmente_vazia(row: tuple) -> bool:
    """True se a linha não tem referência, cliente, especialidade nem valor."""
    return (
        _texto(row[T_REF]) is None
        and _texto(row[T_CLIENTE]) is None
        and _texto(row[T_ESPECIALIDADE]) is None
        and _decimal(row[T_VALOR]) is None
    )


def _separar_ref(referencia: str) -> tuple[int, str] | None:
    """Divide "092.20" em (92, "20"). Devolve None se o formato não bater."""
    if "." not in referencia:
        return None
    num_str, _, ano = referencia.rpartition(".")
    try:
        return int(num_str), ano
    except ValueError:
        return None


def _proxima_referencia_livre(referencia: str, usadas: set[str]) -> str:
    """Devolve o próximo número livre do mesmo ano da referência duplicada.

    Ex.: 092.20 duplicado, com o ano 20 já usado até 179.20 -> 180.20.
    Preserva o número de dígitos (zero-padding) da referência original.
    """
    partes = _separar_ref(referencia)
    if partes is None:
        # Formato inesperado: acrescenta sufixo para garantir unicidade.
        novo = referencia
        i = 2
        while novo in usadas:
            novo = f"{referencia}-{i}"
            i += 1
        return novo

    _, ano = partes
    largura = len(referencia.rpartition(".")[0])

    # Maior número já usado nesse ano.
    maior = 0
    for r in usadas:
        p = _separar_ref(r)
        if p is not None and p[1] == ano:
            maior = max(maior, p[0])

    proximo = maior + 1
    candidato = f"{proximo:0{largura}d}.{ano}"
    while candidato in usadas:
        proximo += 1
        candidato = f"{proximo:0{largura}d}.{ano}"
    return candidato


# ---- Apoio ----
def _get_or_create_especialidade(db, codigo, prazo_top, prazo_ok, rel, cache) -> Especialidade:
    if codigo in cache:
        return cache[codigo]
    obj = db.scalar(select(Especialidade).where(Especialidade.codigo == codigo))
    if not obj:
        obj = Especialidade(codigo=codigo, prazo_top=prazo_top, prazo_ok=prazo_ok)
        db.add(obj)
        db.flush()
        rel.especialidades += 1
    cache[codigo] = obj
    return obj


def _get_or_create_intermediario(db, nome, valor, rel, cache) -> Intermediario:
    if nome in cache:
        return cache[nome]
    obj = db.scalar(select(Intermediario).where(Intermediario.nome == nome))
    if not obj:
        obj = Intermediario(nome=nome, valor=_decimal(valor))
        db.add(obj)
        db.flush()
        rel.intermediarios += 1
    cache[nome] = obj
    return obj


def _get_or_create_externo(db, nome, rel, cache) -> Externo:
    if nome in cache:
        return cache[nome]
    obj = db.scalar(select(Externo).where(Externo.nome == nome))
    if not obj:
        obj = Externo(nome=nome)
        db.add(obj)
        db.flush()
        rel.externos += 1
    cache[nome] = obj
    return obj


def importar_apoio(db: Session, wb, rel: Relatorio) -> tuple[dict, dict, dict]:
    ws = wb["Dados"]
    rows = list(ws.iter_rows(min_row=1, values_only=True))

    espec_map: dict[str, Especialidade] = {}
    interm_map: dict[str, Intermediario] = {}
    ext_map: dict[str, Externo] = {}

    for idx, row in enumerate(rows, start=1):
        if idx >= DADOS_ESPEC_PRIMEIRA_LINHA and len(row) > DADOS_ESPEC_CODIGO:
            codigo = _texto(row[DADOS_ESPEC_CODIGO])
            if codigo and codigo != "ESPEC.":
                top = _inteiro(row[DADOS_ESPEC_PRAZO_TOP]) if len(row) > DADOS_ESPEC_PRAZO_TOP else 0
                ok = _inteiro(row[DADOS_ESPEC_PRAZO_OK]) if len(row) > DADOS_ESPEC_PRAZO_OK else 0
                _get_or_create_especialidade(db, codigo, top, ok, rel, espec_map)

        if idx >= 2 and len(row) > DADOS_INTERM_VALOR:
            nome = _texto(row[DADOS_INTERM_NOME])
            if nome and nome != "INTERMEDIÁRIOS":
                _get_or_create_intermediario(db, nome, row[DADOS_INTERM_VALOR], rel, interm_map)

        if idx >= 2 and len(row) > DADOS_EXTERNO_NOME:
            nome = _texto(row[DADOS_EXTERNO_NOME])
            if nome and nome != "EXTERNOS":
                _get_or_create_externo(db, nome, rel, ext_map)

    db.commit()
    return espec_map, interm_map, ext_map


def importar_equipamentos(db: Session, wb, rel: Relatorio) -> None:
    """Importa os equipamentos (para depreciação) da aba 'Relatório Contas'.

    Idempotente: um equipamento com o mesmo nome não é duplicado.
    A percentagem de alguns blocos é herdada do bloco anterior (col_pct=None).
    """
    if "Relatório Contas" not in wb.sheetnames:
        return
    ws = wb["Relatório Contas"]
    rows = list(ws.iter_rows(min_row=1, values_only=True))

    def celula(linha_idx: int, col: int):
        linha = rows[linha_idx - 1] if len(rows) >= linha_idx else ()
        return linha[col] if len(linha) > col else None

    ultima_pct = None
    for bloco in RC_EQUIPAMENTOS:
        nome = _texto(celula(RC_EQUIP_LINHA_NOME, bloco["col_nome"]))
        if not nome:
            continue

        # Percentagem: própria do bloco ou herdada do anterior.
        if bloco["col_pct"] is not None:
            pct = _decimal(celula(RC_EQUIP_LINHA_DADOS, bloco["col_pct"]))
            if pct is not None:
                ultima_pct = pct
        pct = ultima_pct if ultima_pct is not None else Decimal("0")

        ano = _inteiro(celula(RC_EQUIP_LINHA_DADOS, bloco["col_ano"]))
        valor = _decimal(celula(RC_EQUIP_LINHA_DADOS, bloco["col_valor"]))
        if not ano or valor is None:
            continue

        # Idempotência por nome.
        existente = db.scalar(select(Equipamento).where(Equipamento.nome == nome))
        if existente:
            continue

        db.add(
            Equipamento(nome=nome, percentagem=pct, ano_aquisicao=ano, valor=valor)
        )
        rel.equipamentos += 1

    db.commit()


def _resolver_especialidade(db, codigo, espec_map, rel) -> int | None:
    if not codigo:
        return None
    esp = espec_map.get(codigo)
    if not esp:
        esp = _get_or_create_especialidade(db, codigo, 0, 0, rel, espec_map)
        rel.nomes_nao_resolvidos.add(f"especialidade:{codigo}")
    return esp.id


def _resolver_externo(db, nome, ext_map, rel) -> int | None:
    if not nome:
        return None
    if nome not in ext_map:
        rel.nomes_nao_resolvidos.add(f"externo:{nome}")
    ex = _get_or_create_externo(db, nome, rel, ext_map)
    return ex.id


def _resolver_intermediario(db, nome, interm_map, rel) -> int | None:
    if not nome:
        return None
    if nome not in interm_map:
        rel.nomes_nao_resolvidos.add(f"intermediario:{nome}")
    it = _get_or_create_intermediario(db, nome, None, rel, interm_map)
    return it.id


# A situação vem da COR de fundo da coluna A (PONT. SITUAÇÃO):
#   verde  (92D050) -> Pago
#   laranja (FFC000) -> Falta pagamento
#   amarelo (FFFF00) -> Em curso
# O texto é secundário; usa-se apenas se não houver cor reconhecida.
_COR_SITUACAO = {
    "FF92D050": Situacao.PAGO,
    "FFFFC000": Situacao.FALTA_PAGAMENTO,
    "FFFFFF00": Situacao.EM_CURSO,
}
_TEXTO_SITUACAO = {
    "pago": Situacao.PAGO,
    "falta pagamento": Situacao.FALTA_PAGAMENTO,
    "em curso": Situacao.EM_CURSO,
}


def _situacao_da_celula(cel) -> str | None:
    """Determina a situação a partir da cor de fundo (prioritário) ou do texto."""
    # Cor de fundo.
    fill = getattr(cel, "fill", None)
    if fill is not None and getattr(fill, "patternType", None):
        rgb = getattr(fill.fgColor, "rgb", None)
        if isinstance(rgb, str):
            sit = _COR_SITUACAO.get(rgb.upper())
            if sit:
                return sit.value
    # Texto (secundário).
    texto = _texto(cel.value)
    if texto:
        sit = _TEXTO_SITUACAO.get(texto.lower())
        if sit:
            return sit.value
    return None


def _construir_item(db, row, situacao, espec_map, ext_map, rel) -> TrabalhoItem:
    return TrabalhoItem(
        especialidade_id=_resolver_especialidade(db, _texto(row[T_ESPECIALIDADE]), espec_map, rel),
        externo_id=_resolver_externo(db, _texto(row[T_EXTERNO]), ext_map, rel),
        valor=_decimal(row[T_VALOR]),
        valor_externo=_decimal(row[T_VALOR_EXTERNO]),
        data_adjudicacao=_data(row[T_ADJUDICACAO]),
        data_entrega=_data(row[T_ENTREGA]),
        situacao=situacao,
        observacoes=_texto(row[T_OBSERVACOES]),
    )


def importar_trabalhos(db, wb, espec_map, interm_map, ext_map, rel: Relatorio) -> None:
    ws = wb["Trabalhos"]

    # Referências que já estavam na BD antes desta importação (reimportação -> ignorar).
    pre_existentes = {
        r for (r,) in db.execute(select(Trabalho.referencia).where(Trabalho.referencia.is_not(None)))
    }
    # Todas as referências conhecidas (BD + criadas nesta execução), para detetar
    # duplicados dentro do próprio ficheiro e calcular o próximo número livre.
    usadas = set(pre_existentes)

    trabalho_atual: Trabalho | None = None
    ignorar_atual = False  # quando o trabalho corrente já existia na BD
    criados: list[Trabalho] = []  # trabalhos criados nesta execução (para herança da situação)

    # Itera as células (não values_only) para poder ler a COR de fundo da coluna A.
    for celulas in ws.iter_rows(min_row=2):
        if len(celulas) <= T_OBSERVACOES:
            continue

        row = tuple(c.value for c in celulas)
        situacao = _situacao_da_celula(celulas[T_SITUACAO])

        if _linha_totalmente_vazia(row):
            rel.linhas_vazias += 1
            continue

        referencia = _texto(row[T_REF])
        cliente = _texto(row[T_CLIENTE])
        inicia_trabalho = referencia is not None or cliente is not None

        if inicia_trabalho:
            # Reimportação: se a referência já estava na BD antes desta execução,
            # ignora este trabalho e as suas linhas de continuação.
            if referencia and referencia in pre_existentes:
                trabalho_atual = None
                ignorar_atual = True
                rel.trabalhos_ignorados += 1
                continue

            # Duplicado DENTRO do ficheiro: a referência já foi usada nesta execução.
            # Atribui o próximo número livre do mesmo ano (não perde o trabalho).
            if referencia and referencia in usadas:
                nova = _proxima_referencia_livre(referencia, usadas)
                rel.referencias_renumeradas.append((referencia, nova))
                referencia = nova

            ignorar_atual = False
            trabalho_atual = Trabalho(
                referencia=referencia,
                cliente=cliente,
                contacto=_texto(row[T_CONTACTO]),
                localidade=_texto(row[T_LOCALIDADE]),
                endereco=_texto(row[T_ENDERECO]),
                intermediario_id=_resolver_intermediario(
                    db, _texto(row[T_INTERMEDIARIO]), interm_map, rel
                ),
            )
            db.add(trabalho_atual)
            criados.append(trabalho_atual)
            rel.trabalhos_criados += 1
            if referencia:
                usadas.add(referencia)

            # A própria linha inicial é também um item (tem especialidade/valor).
            if _texto(row[T_ESPECIALIDADE]) or _decimal(row[T_VALOR]) or _data(row[T_ADJUDICACAO]):
                trabalho_atual.itens.append(
                    _construir_item(db, row, situacao, espec_map, ext_map, rel)
                )
                rel.itens_criados += 1
        else:
            # Linha de continuação: item adicional do trabalho corrente.
            if ignorar_atual or trabalho_atual is None:
                continue
            trabalho_atual.itens.append(
                _construir_item(db, row, situacao, espec_map, ext_map, rel)
            )
            rel.itens_criados += 1

        if rel.trabalhos_criados % 200 == 0:
            db.flush()

    # Herança da situação: itens sem situação herdam a PRIMEIRA situação preenchida
    # que exista no trabalho (mesmo que não seja no primeiro item).
    for trab in criados:
        situacao_base = next((i.situacao for i in trab.itens if i.situacao), None)
        if not situacao_base:
            continue
        for item in trab.itens:
            if not item.situacao:
                item.situacao = situacao_base

    db.commit()


def importar(path: str) -> Relatorio:
    rel = Relatorio()
    # Sem read_only: é necessário para aceder à cor de fundo das células (situação).
    wb = load_workbook(path, data_only=True)
    db = SessionLocal()
    try:
        espec_map, interm_map, ext_map = importar_apoio(db, wb, rel)
        importar_equipamentos(db, wb, rel)
        importar_trabalhos(db, wb, espec_map, interm_map, ext_map, rel)
    finally:
        db.close()
        wb.close()
    return rel


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print('Uso: python -m app.importar "caminho/para/ficheiro.xlsm"')
        sys.exit(1)
    relatorio = importar(sys.argv[1])
    relatorio.imprimir()
